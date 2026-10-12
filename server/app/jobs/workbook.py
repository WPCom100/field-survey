"""Read a master / floor survey workbook into plain rows (PLAN §7.5, T27).

Nothing here knows about the database: the reader returns strings, numbers and dates per tab,
the Summary's floor list with its Surveyed/Updated text dates, and the Photo Log's embedded
images. The mapping to records is `app.jobs.seed_import`.

Tabs are read by header name (row 3, the master's names), so a floor workbook with the same
headers imports the same way; `ColumnMap` lets the column-mapping UI (T28) rename headers per
tab. A tab that is missing is reported in `checks` and skipped, never silently passed.

The Photo Log images are not cells: they are `twoCellAnchor` pictures in the sheet's drawing,
anchored by row. openpyxl only loads pictures when Pillow is installed, so the drawing XML is
read from the zip directly (sheet → drawing → media), which also keeps the exact JPEG bytes.
"""

from __future__ import annotations

import hashlib
import re
import struct
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import IO, Any
from xml.etree import ElementTree as ET

import openpyxl

ColumnMap = dict[str, dict[str, str]]
"""{tab: {field: header}} overrides for workbooks whose headers differ from the master."""

HEADER_ROW = 3
FIRST_DATA_ROW = 4

# field → master header, per tab (PLAN §2.2 bold columns; D76 names on our side)
DEVICE_COLUMNS = {
    "level": "Level",
    "display_id": "ID",
    "category": "Category",
    "device": "Device",
    "type": "Type",
    "room_location": "Room / Location",
    "drawing_position": "Drawing Position",
    "status": "Status",
    "notes": "Notes",
    "photo": "Photo",
}
CONNECTION_COLUMNS = {
    "level": "Level",
    "display_id": "ID",
    "from_text": "From",
    "cable": "Cable",
    "to_text": "To",
    "qty": "Qty",
    "status": "Status",
    "note": "Note",
}
PUNCH_COLUMNS = {
    "level": "Level",
    "number": "#",
    "device": "Device",
    "item": "Item",
    "type": "Type",
    "photo": "Photo",
    "owner": "Owner",
    "tracking_status": "Tracking status",
    "date_closed": "Date closed",
    "notes": "Notes",
}
PART_COLUMNS = {
    "level": "Level",
    "number": "#",
    "device_location": "Device / Location",
    "part": "Part",
    "part_no": "Part no.",
    "qty": "Qty",
    "owner": "Owner",
    "source": "Source",
    "part_status": "Part status",
    "notes": "Notes",
}
PHOTO_COLUMNS = {
    "level": "Level",
    "number": "Photo",
    "location": "Location",
    "subject": "Subject",
    "what_it_shows": "What it shows",
    "device_id": "Device ID",
}
CHANGE_COLUMNS = {
    "date": "Date",
    "level": "Level",
    "sheet": "Sheet",
    "ref": "Ref",
    "entry": "Entry",
    "field": "Field",
    "was": "Was",
    "now": "Now",
    "change": "Change",
    "reason": "Reason",
}
TABS: dict[str, dict[str, str]] = {
    "Devices": DEVICE_COLUMNS,
    "Connections": CONNECTION_COLUMNS,
    "Punch List": PUNCH_COLUMNS,
    "Missing Parts": PART_COLUMNS,
    "Photo Log": PHOTO_COLUMNS,
    "Change Log": CHANGE_COLUMNS,
}
# Tabs a workbook must have to be a seed at all; the others are optional (a floor workbook has
# no Change Log, a fresh master no Photo Log).
REQUIRED_TABS = ("Devices",)

SUMMARY_TAB = "Summary"
# Summary rows 8–15 carry Floor / Surveyed / Updated (T02 follow-up): found by the header row
# whose first cell is "Floor", read until "Total" or a blank floor.
SUMMARY_FLOOR_HEADER = "Floor"


@dataclass
class Row:
    """One data row: `values` by field name, plus where it came from for messages."""

    tab: str
    excel_row: int
    values: dict[str, Any]

    def get(self, name: str) -> Any:
        return self.values.get(name)


@dataclass
class LevelRow:
    code: str
    floor_order: int
    surveyed: date | None
    updated: date | None


@dataclass
class Image:
    data: bytes
    sha256: str
    mime: str
    width: int | None
    height: int | None


@dataclass
class Workbook:
    name: str
    sha256: str
    levels: list[LevelRow]
    rows: dict[str, list[Row]]  # tab → rows
    images: dict[int, Image]  # Photo Log excel row → image
    checks: list[str] = field(default_factory=list)  # problems a person should read

    def tab(self, name: str) -> list[Row]:
        return self.rows.get(name, [])


class WorkbookError(Exception):
    pass


# ------------------------------------------------------------------ cell values


def clean(value: Any) -> Any:
    """Strings stripped and emptied to None; dates kept; Excel's numeric text left as text."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def as_text(value: Any) -> str | None:
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value)


def as_int(value: Any) -> int | None:
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, bool):
        raise WorkbookError(f"not a number: {value!r}")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip().lstrip("#")
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    raise WorkbookError(f"not a number: {value!r}")


def as_date(value: Any) -> date | None:
    """Excel dates arrive as datetimes; the Summary keeps ISO text (T02 follow-up)."""
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value)
    try:
        return date.fromisoformat(text[:10])
    except ValueError as e:
        raise WorkbookError(f"not a date: {value!r}") from e


# --------------------------------------------------------------------- reading


def read_workbook(
    source: str | Path | IO[bytes], name: str | None = None, mapping: ColumnMap | None = None
) -> Workbook:
    """Read every known tab. `source` is a path or a binary file object."""
    if isinstance(source, (str, Path)):
        path = Path(source)
        data = path.read_bytes()
        name = name or path.name
    else:
        data = source.read()
        name = name or "workbook.xlsx"
    sha = hashlib.sha256(data).hexdigest()
    checks: list[str] = []
    import io

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception as e:  # zip or xml trouble
        raise WorkbookError(f"{name}: not an .xlsx workbook ({e})") from e

    rows: dict[str, list[Row]] = {}
    for tab, columns in TABS.items():
        if tab not in wb.sheetnames:
            if tab in REQUIRED_TABS:
                raise WorkbookError(f"{name}: tab {tab!r} is missing")
            checks.append(f"tab {tab!r} missing; skipped")
            continue
        rows[tab] = _read_tab(wb[tab], tab, columns, (mapping or {}).get(tab, {}), checks)

    levels = _read_summary(wb, checks) if SUMMARY_TAB in wb.sheetnames else []
    if not levels:
        levels = _levels_from_rows(rows)
        checks.append("no Summary floor table; levels taken from the data tabs in order")

    images = _read_images(io.BytesIO(data), "Photo Log", checks) if "Photo Log" in rows else {}
    return Workbook(name=name, sha256=sha, levels=levels, rows=rows, images=images, checks=checks)


def _header_index(ws, tab: str, columns: dict[str, str], override: dict[str, str], checks):
    headers: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h = as_text(ws.cell(HEADER_ROW, c).value)
        if h:
            headers.setdefault(h.casefold(), c)
    index: dict[str, int] = {}
    for fld, header in columns.items():
        wanted = override.get(fld, header)
        col = headers.get(wanted.casefold())
        if col is None:
            checks.append(f"{tab}: column {wanted!r} not found (field {fld})")
        else:
            index[fld] = col
    return index


def _read_tab(ws, tab: str, columns, override, checks) -> list[Row]:
    index = _header_index(ws, tab, columns, override, checks)
    out: list[Row] = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        raw = {fld: ws.cell(r, col).value for fld, col in index.items()}
        if all(clean(v) is None for v in raw.values()):
            continue
        values: dict[str, Any] = {}
        for fld, v in raw.items():
            try:
                values[fld] = _convert(tab, fld, v)
            except WorkbookError as e:
                checks.append(f"{tab} row {r}: {fld}: {e}")
                values[fld] = None
        out.append(Row(tab=tab, excel_row=r, values=values))
    return out


INT_FIELDS = {
    ("Connections", "qty"),
    ("Punch List", "number"),
    ("Missing Parts", "number"),
    ("Missing Parts", "qty"),
    ("Photo Log", "number"),
}
DATE_FIELDS = {("Punch List", "date_closed"), ("Change Log", "date")}


def _convert(tab: str, fld: str, value: Any) -> Any:
    if (tab, fld) in INT_FIELDS:
        return as_int(value)
    if (tab, fld) in DATE_FIELDS:
        return as_date(value)
    return as_text(value)


def _read_summary(wb, checks) -> list[LevelRow]:
    ws = wb[SUMMARY_TAB]
    start = None
    for r in range(1, min(ws.max_row, 40) + 1):
        if as_text(ws.cell(r, 1).value) == SUMMARY_FLOOR_HEADER:
            start = r + 1
            break
    if start is None:
        checks.append("Summary: no 'Floor' table found")
        return []
    levels: list[LevelRow] = []
    for r in range(start, ws.max_row + 1):
        code = as_text(ws.cell(r, 1).value)
        if code is None or code.casefold() == "total":
            break
        try:
            surveyed = as_date(ws.cell(r, 2).value)
            updated = as_date(ws.cell(r, 3).value)
        except WorkbookError as e:
            checks.append(f"Summary row {r}: {e}")
            surveyed = updated = None
        levels.append(
            LevelRow(code=code, floor_order=len(levels) + 1, surveyed=surveyed, updated=updated)
        )
    return levels


def _levels_from_rows(rows: dict[str, list[Row]]) -> list[LevelRow]:
    seen: list[str] = []
    for tab in ("Devices", "Connections", "Punch List", "Missing Parts", "Photo Log"):
        for row in rows.get(tab, []):
            code = row.get("level")
            if code and code not in seen:
                seen.append(code)
    return [
        LevelRow(code=c, floor_order=i + 1, surveyed=None, updated=None) for i, c in enumerate(seen)
    ]


# ---------------------------------------------------------------------- images

NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "wb": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}
MIME_BY_EXT = {".jpeg": "image/jpeg", ".jpg": "image/jpeg", ".png": "image/png"}


def _rels(z: zipfile.ZipFile, part: str) -> dict[str, str]:
    """rId → absolute part name for `part`'s relationships file."""
    folder, _, base = part.rpartition("/")
    rels_name = f"{folder}/_rels/{base}.rels"
    if rels_name not in z.namelist():
        return {}
    out: dict[str, str] = {}
    for rel in ET.fromstring(z.read(rels_name)).findall("rel:Relationship", NS):
        target = rel.get("Target", "")
        if not target.startswith("/"):
            target = f"{folder}/{target}"
        out[rel.get("Id", "")] = _normpath(target.lstrip("/"))
    return out


def _normpath(p: str) -> str:
    parts: list[str] = []
    for seg in p.split("/"):
        if seg == "..":
            if parts:
                parts.pop()
        elif seg and seg != ".":
            parts.append(seg)
    return "/".join(parts)


def _read_images(fileobj: IO[bytes], tab: str, checks) -> dict[int, Image]:
    """Pictures anchored in `tab`, by 1-based Excel row of their top-left anchor."""
    out: dict[int, Image] = {}
    with zipfile.ZipFile(fileobj) as z:
        wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
        wb_rels = _rels(z, "xl/workbook.xml")
        sheet_part = None
        sheets = wb_xml.find("wb:sheets", NS)
        for sheet in sheets if sheets is not None else []:
            if sheet.get("name") == tab:
                rid = sheet.get(f"{{{NS['r']}}}id")
                sheet_part = wb_rels.get(rid or "")
        if sheet_part is None or sheet_part not in z.namelist():
            return out
        sheet_xml = ET.fromstring(z.read(sheet_part))
        drawing_el = sheet_xml.find("wb:drawing", NS)
        if drawing_el is None:
            return out
        drawing_part = _rels(z, sheet_part).get(drawing_el.get(f"{{{NS['r']}}}id") or "")
        if not drawing_part or drawing_part not in z.namelist():
            checks.append(f"{tab}: drawing part missing; no images")
            return out
        media = _rels(z, drawing_part)
        for anchor in ET.fromstring(z.read(drawing_part)):
            frm = anchor.find("xdr:from", NS)
            blip = anchor.find(".//a:blip", NS)
            if frm is None or blip is None:
                continue
            row = int(frm.findtext("xdr:row", default="0", namespaces=NS)) + 1
            part = media.get(blip.get(f"{{{NS['r']}}}embed") or "")
            if not part or part not in z.namelist():
                checks.append(f"{tab} row {row}: picture without media")
                continue
            data = z.read(part)
            ext = Path(part).suffix.lower()
            mime = MIME_BY_EXT.get(ext, "application/octet-stream")
            size = jpeg_size(data) if mime == "image/jpeg" else png_size(data)
            if row in out:
                checks.append(f"{tab} row {row}: two pictures; the first is kept")
                continue
            out[row] = Image(
                data=data,
                sha256=hashlib.sha256(data).hexdigest(),
                mime=mime,
                width=size[0] if size else None,
                height=size[1] if size else None,
            )
    return out


def jpeg_size(data: bytes) -> tuple[int, int] | None:
    """Width/height from the first SOF marker; None when the stream has none (not a JPEG)."""
    if data[:2] != b"\xff\xd8":
        return None
    i = 2
    n = len(data)
    while i + 9 < n:
        if data[i] != 0xFF:
            i += 1
            continue
        marker = data[i + 1]
        if marker in (0xD8, 0x01) or 0xD0 <= marker <= 0xD7:
            i += 2
            continue
        length = struct.unpack(">H", data[i + 2 : i + 4])[0]
        if marker in (0xC0, 0xC1, 0xC2, 0xC3, 0xC5, 0xC6, 0xC7, 0xC9, 0xCA, 0xCB, 0xCD, 0xCE, 0xCF):
            height, width = struct.unpack(">HH", data[i + 5 : i + 9])
            return width, height
        i += 2 + length
    return None


def png_size(data: bytes) -> tuple[int, int] | None:
    if data[:8] != b"\x89PNG\r\n\x1a\n" or len(data) < 24:
        return None
    width, height = struct.unpack(">II", data[16:24])
    return width, height
