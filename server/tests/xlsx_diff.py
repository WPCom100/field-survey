"""TP-4: diff an exported workbook against the master, cell by cell (T29, PLAN §7.6).

    python tests/xlsx_diff.py fixtures/Lighting_Control_Survey_Master_B2-06_4.xlsx export.xlsx

Compares, per tab: every cell value (formulas as text) and its style (font, fill, number format,
alignment, borders), column widths, row heights (reported, not failed), freeze panes, Tables
(name, range, style, autofilter, columns), data validations, conditional formats, merged ranges,
page setup, and the drawings — charts (type, grouping, series refs and colours, anchor and size)
and pictures (anchor, size and the media bytes' SHA-256).

Every difference is listed; the ones the export is *meant* to produce are tagged with a reason
(`Difference.reason`) and `DiffResult.unexpected()` is what must be empty:

| Reason | What |
|---|---|
| D51 | no Change Log tab |
| D50 | a Unison bus spelling in Connections!From exported as `Stations & Sensors` |
| D110 / D138 | an empty device or connection status exported as `Unsurveyed` |
| photo list spacing | `12,13` in the master, `12, 13` in the export |
| device ID range | `01-S8 to 01-S12` in the master, the IDs listed in the export |
| Source parenthetical | master `Connections C-06 (fiber to Level 01)`, export `Connections C-06` |
| allow | a cell the caller lists in `allow` = {(tab, header, (Level, ID or #)): reason} |
"""

from __future__ import annotations

import hashlib
import re
import sys
import zipfile
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter

DATA_TABS = ("Devices", "Connections", "Punch List", "Missing Parts", "Photo Log")
BUS_LABEL = "Stations & Sensors"
UNSURVEYED = "Unsurveyed"
EMU_PER_PX = 9525
BOX_TOLERANCE_PX = 8
NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "wb": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}


@dataclass
class Difference:
    tab: str
    where: str
    kind: str
    master: Any
    export: Any
    reason: str | None = None

    def __str__(self) -> str:
        tag = f" [{self.reason}]" if self.reason else ""
        return (
            f"{self.tab}!{self.where} {self.kind}: "
            f"master={self.master!r} export={self.export!r}{tag}"
        )


@dataclass
class DiffResult:
    differences: list[Difference] = field(default_factory=list)
    row_heights: dict[str, tuple[int, int]] = field(default_factory=dict)  # tab → (off, rows)

    def unexpected(self) -> list[Difference]:
        return [d for d in self.differences if d.reason is None]

    def documented(self) -> Counter:
        return Counter(d.reason for d in self.differences if d.reason)

    def report(self, limit: int = 60) -> str:
        lines = [
            f"{len(self.unexpected())} unexpected, {sum(self.documented().values())} documented"
        ]
        for d in self.unexpected()[:limit]:
            lines.append("  " + str(d))
        if len(self.unexpected()) > limit:
            lines.append(f"  … {len(self.unexpected()) - limit} more")
        for reason, n in sorted(self.documented().items()):
            lines.append(f"  documented {reason}: {n}")
        for tab, (off, rows) in self.row_heights.items():
            lines.append(f"  row heights {tab}: {off} of {rows} differ")
        return "\n".join(lines)


# ----------------------------------------------------------------- helpers


def _value(v: Any) -> Any:
    if v is None or v == "":
        return None
    if isinstance(v, datetime) and (v.hour, v.minute, v.second) == (0, 0, 0):
        return v.date()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _rgb(color) -> str | None:
    if color is None:
        return None
    rgb = getattr(color, "rgb", None)
    return rgb if isinstance(rgb, str) else None


def _style(cell) -> tuple:
    f = cell.font
    fill = cell.fill
    fill_rgb = _rgb(fill.fgColor) if fill is not None and fill.fill_type else None
    if fill_rgb in (None, "00000000") and fill is not None and fill.fill_type:
        fill_rgb = _rgb(fill.bgColor)
    b = cell.border
    border = tuple(getattr(b, side).style for side in ("left", "right", "top", "bottom"))
    al = cell.alignment
    return (
        f.name,
        float(f.sz) if f.sz else None,
        bool(f.b),
        _rgb(f.color),
        fill_rgb,
        cell.number_format,
        bool(al.wrap_text),
        al.horizontal,
        al.vertical,
        border,
    )


def _widths(ws, ncols: int) -> dict[int, float | None]:
    out: dict[int, float | None] = dict.fromkeys(range(1, ncols + 1))
    for dim in ws.column_dimensions.values():
        if not dim.width:
            continue
        lo = dim.min or openpyxl.utils.column_index_from_string(dim.index)
        hi = dim.max or lo
        for c in range(lo, min(hi, ncols) + 1):
            out[c] = round(float(dim.width), 2)
    return out


def _col_px(ws, c0: int) -> int:
    """Excel's pixel width of a zero-based column (Calibri 11: 7 px per unit)."""
    for dim in ws.column_dimensions.values():
        if not dim.width:
            continue
        lo = (dim.min or openpyxl.utils.column_index_from_string(dim.index)) - 1
        hi = (dim.max or lo + 1) - 1
        if lo <= c0 <= hi:
            # the stored width already carries Excel's 5-px padding: px = width × 7
            return round(float(dim.width) * 7)
    return 64


def _row_px(ws, r0: int) -> int:
    h = ws.row_dimensions[r0 + 1].height if (r0 + 1) in ws.row_dimensions else None
    return int(4.0 / 3.0 * (h if h else 15))


def _rels(z: zipfile.ZipFile, part: str) -> dict[str, str]:
    folder, _, base = part.rpartition("/")
    name = f"{folder}/_rels/{base}.rels"
    if name not in z.namelist():
        return {}
    out = {}
    for rel in ET.fromstring(z.read(name)).findall("rel:Relationship", NS):
        target = rel.get("Target", "")
        if not target.startswith("/"):
            target = f"{folder}/{target}"
        parts: list[str] = []
        for seg in target.lstrip("/").split("/"):
            if seg == "..":
                parts.pop()
            elif seg and seg != ".":
                parts.append(seg)
        out[rel.get("Id", "")] = "/".join(parts)
    return out


def _drawings(path: Path, wb) -> dict[str, list[dict[str, Any]]]:
    """Per sheet: the drawing anchors as pixel boxes — {kind, col, row, x, y, w, h, sha}."""
    out: dict[str, list[dict[str, Any]]] = {}
    with zipfile.ZipFile(path) as z:
        wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
        wb_rels = _rels(z, "xl/workbook.xml")
        sheets = wb_xml.find("wb:sheets", NS)
        for sheet in sheets if sheets is not None else []:
            name = sheet.get("name")
            part = wb_rels.get(sheet.get(f"{{{NS['r']}}}id") or "")
            if not part or part not in z.namelist() or name not in wb.sheetnames:
                continue
            ws = wb[name]
            sx = ET.fromstring(z.read(part))
            d = sx.find("wb:drawing", NS)
            if d is None:
                out[name] = []
                continue
            dpart = _rels(z, part).get(d.get(f"{{{NS['r']}}}id") or "")
            if not dpart or dpart not in z.namelist():
                out[name] = []
                continue
            media = _rels(z, dpart)
            anchors = []
            for anchor in ET.fromstring(z.read(dpart)):
                frm = anchor.find("xdr:from", NS)
                if frm is None:
                    continue
                col = int(frm.findtext("xdr:col", "0", NS))
                row = int(frm.findtext("xdr:row", "0", NS))
                x = (
                    sum(_col_px(ws, c) for c in range(col))
                    + int(frm.findtext("xdr:colOff", "0", NS)) / EMU_PER_PX
                )
                y = (
                    sum(_row_px(ws, r) for r in range(row))
                    + int(frm.findtext("xdr:rowOff", "0", NS)) / EMU_PER_PX
                )
                ext = anchor.find("xdr:ext", NS)
                to = anchor.find("xdr:to", NS)
                if ext is not None:
                    w = int(ext.get("cx", "0")) / EMU_PER_PX
                    h = int(ext.get("cy", "0")) / EMU_PER_PX
                elif to is not None:
                    tcol = int(to.findtext("xdr:col", "0", NS))
                    trow = int(to.findtext("xdr:row", "0", NS))
                    x2 = (
                        sum(_col_px(ws, c) for c in range(tcol))
                        + int(to.findtext("xdr:colOff", "0", NS)) / EMU_PER_PX
                    )
                    y2 = (
                        sum(_row_px(ws, r) for r in range(trow))
                        + int(to.findtext("xdr:rowOff", "0", NS)) / EMU_PER_PX
                    )
                    w, h = x2 - x, y2 - y
                else:
                    w = h = 0.0
                pic = anchor.find("xdr:pic", NS)
                kind, sha = ("chart", None)
                if pic is not None:
                    blip = pic.find(".//a:blip", NS)
                    mpart = (
                        media.get(blip.get(f"{{{NS['r']}}}embed") or "")
                        if blip is not None
                        else None
                    )
                    sha = (
                        hashlib.sha256(z.read(mpart)).hexdigest()
                        if mpart and mpart in z.namelist()
                        else None
                    )
                    kind = "pic"
                anchors.append(
                    {
                        "kind": kind,
                        "col": col,
                        "row": row,
                        "x": x,
                        "y": y,
                        "w": w,
                        "h": h,
                        "sha": sha,
                    }
                )
            out[name] = anchors
    return out


def _ref(obj) -> str | None:
    if obj is None:
        return None
    for attr in ("numRef", "strRef"):
        ref = getattr(obj, attr, None)
        if ref is not None and ref.f:
            return ref.f.replace("$", "")
    return None


def _chart_shape(chart) -> dict[str, Any]:
    series = []
    for s in chart.series:
        fill = None
        gp = s.graphicalProperties
        if gp is not None and gp.solidFill is not None and gp.solidFill.srgbClr is not None:
            clr = gp.solidFill.srgbClr
            fill = (clr if isinstance(clr, str) else clr.val).upper()
        series.append({"name": _ref(s.tx), "cat": _ref(s.cat), "val": _ref(s.val), "fill": fill})
    return {
        "type": type(chart).__name__,
        "grouping": chart.grouping,
        "bar": getattr(chart, "type", None),
        "overlap": getattr(chart, "overlap", None),
        "gap": getattr(chart, "gapWidth", None),
        "legend": chart.legend.position if chart.legend else None,
        "title": None if chart.title is None else "set",
        "series": series,
    }


def _bus_spelling(text: str | None) -> bool:
    if not text:
        return False
    t = text.casefold()
    return "station" in t and ("sensor" in t or t.endswith("stations"))


def _expand_range(text: str) -> str:
    out = []
    for part in (p.strip() for p in text.split(",")):
        ends = [e.strip() for e in re.split(r"\bto\b", part) if e.strip()]
        if len(ends) == 2:
            a, b = (re.fullmatch(r"(.*?)(\d+)", e) for e in ends)
            if a and b and a.group(1) == b.group(1):
                lo, hi = sorted((int(a.group(2)), int(b.group(2))))
                out.extend(f"{a.group(1)}{n}" for n in range(lo, hi + 1))
                continue
        out.append(part)
    return ", ".join(out)


def classify(tab: str, header: str | None, key: tuple, a: Any, b: Any, allow: dict) -> str | None:
    """The documented reason for a value difference, or None."""
    if (tab, header, key) in allow:
        return allow[(tab, header, key)]
    if tab == "Connections" and header == "From" and _bus_spelling(a) and b == BUS_LABEL:
        return "D50"
    if header == "Status" and a is None and b == UNSURVEYED:
        return "D110" if tab == "Devices" else "D138"
    if (
        header == "Photo"
        and isinstance(a, str)
        and isinstance(b, str)
        and re.sub(r",\s*", ", ", a) == b
    ):
        return "photo list spacing"
    if header == "Device ID" and isinstance(a, str) and " to " in a and _expand_range(a) == b:
        return "device ID range"
    if header == "Source" and isinstance(a, str) and re.sub(r"\s*\([^)]*\)", "", a) == b:
        return "Source parenthetical"
    return None


# -------------------------------------------------------------------- diff


def diff(
    master_path: str | Path,
    export_path: str | Path,
    *,
    ignore_tabs: Iterable[str] = ("Change Log",),
    allow: dict[tuple, str] | None = None,
) -> DiffResult:
    allow = allow or {}
    master_path, export_path = Path(master_path), Path(export_path)
    m = openpyxl.load_workbook(master_path)
    e = openpyxl.load_workbook(export_path)
    res = DiffResult()
    ignore = set(ignore_tabs)

    m_tabs = [s for s in m.sheetnames if s not in ignore]
    if m_tabs != e.sheetnames:
        res.differences.append(
            Difference("(workbook)", "sheets", "tabs", m.sheetnames, e.sheetnames)
        )
    for s in m.sheetnames:
        if s in ignore and s not in e.sheetnames:
            res.differences.append(Difference("(workbook)", s, "tab absent", s, None, "D51"))

    for tab in m_tabs:
        if tab not in e.sheetnames:
            continue
        mw, ew = m[tab], e[tab]

        def add(where, kind, a, b, reason=None, *, _tab=tab):
            res.differences.append(Difference(_tab, where, kind, a, b, reason))

        is_data = tab in DATA_TABS
        if is_data:
            ncols = len(next(iter(mw.tables.values())).tableColumns) if mw.tables else mw.max_column
            headers = {c: _value(mw.cell(3, c).value) for c in range(1, ncols + 1)}
        else:
            ncols = max(mw.max_column, ew.max_column)
            headers = {}
        nrows = max(mw.max_row, ew.max_row)
        if is_data and mw.max_row != ew.max_row:
            add("rows", "data rows", mw.max_row - 3, ew.max_row - 3)

        # values and styles
        for r in range(1, nrows + 1):
            key = (
                (str(_value(mw.cell(r, 1).value)), str(_value(mw.cell(r, 2).value)))
                if is_data
                else ()
            )
            for c in range(1, ncols + 1):
                mc, ec = mw.cell(r, c), ew.cell(r, c)
                a, b = _value(mc.value), _value(ec.value)
                where = f"{get_column_letter(c)}{r}"
                header = headers.get(c) if is_data and r > 3 else None
                if a != b:
                    add(where, "value", a, b, classify(tab, header, key, a, b, allow))
                styled = mc.value is not None or mc.border.left.style is not None
                if styled and _style(mc) != _style(ec):
                    add(where, "style", _style(mc), _style(ec))

        # widths, freeze, merged, print
        mwid, ewid = _widths(mw, ncols), _widths(ew, ncols)
        for c in range(1, ncols + 1):
            if mwid[c] is not None and mwid[c] != ewid[c]:
                add(get_column_letter(c), "width", mwid[c], ewid[c])
        mp, ep = mw.sheet_view.pane, ew.sheet_view.pane
        mf = (mp.xSplit, mp.ySplit) if mp is not None else None
        ef = (ep.xSplit, ep.ySplit) if ep is not None else None
        if mf != ef:
            add("pane", "freeze", mf, ef)
        mm, em = sorted(map(str, mw.merged_cells.ranges)), sorted(map(str, ew.merged_cells.ranges))
        if mm != em:
            add("merged", "merged", mm, em)
        for name, mv, ev in (
            ("orientation", mw.page_setup.orientation, ew.page_setup.orientation),
            ("fitToHeight", mw.page_setup.fitToHeight, ew.page_setup.fitToHeight),
            ("fitToWidth", mw.page_setup.fitToWidth, ew.page_setup.fitToWidth),
            (
                "fitToPage",
                mw.sheet_properties.pageSetUpPr.fitToPage
                if mw.sheet_properties.pageSetUpPr
                else None,
                ew.sheet_properties.pageSetUpPr.fitToPage
                if ew.sheet_properties.pageSetUpPr
                else None,
            ),
            ("print_title_rows", mw.print_title_rows, ew.print_title_rows),
        ):
            if mv != ev:
                add("page", name, mv, ev)

        # tables
        def tables(ws):
            out = []
            for t in ws.tables.values():
                si = t.tableStyleInfo
                out.append(
                    (
                        t.name,
                        t.displayName,
                        t.ref,
                        si.name if si else None,
                        bool(si.showRowStripes) if si else None,
                        bool(si.showColumnStripes) if si else None,
                        bool(si.showFirstColumn) if si else None,
                        bool(si.showLastColumn) if si else None,
                        t.autoFilter.ref if t.autoFilter else None,
                        tuple(c.name for c in t.tableColumns),
                    )
                )
            return sorted(out)

        if tables(mw) != tables(ew):
            add("tables", "tables", tables(mw), tables(ew))

        # validations
        def validations(ws):
            return sorted(
                (
                    str(dv.sqref),
                    (dv.formula1 or "").lstrip(
                        "="
                    ),  # openpyxl keeps a leading '=', XlsxWriter drops it
                    dv.type,
                    bool(dv.allow_blank),
                    bool(dv.showDropDown),
                    bool(dv.showErrorMessage),
                    bool(dv.showInputMessage),
                )
                for dv in ws.data_validations.dataValidation
            )

        if validations(mw) != validations(ew):
            add("validations", "validations", validations(mw), validations(ew))

        # conditional formats
        def cfs(ws):
            out = []
            for rng, rules in ws.conditional_formatting._cf_rules.items():
                for rule in rules:
                    d = rule.dxf
                    fill = None
                    if d is not None and d.fill is not None:
                        fill = _rgb(d.fill.bgColor) or _rgb(d.fill.fgColor)
                    out.append(
                        (str(rng.sqref), rule.type, rule.operator, tuple(rule.formula or ()), fill)
                    )
            return sorted(out)

        if cfs(mw) != cfs(ew):
            add("conditional", "conditional formats", cfs(mw), cfs(ew))

        # charts
        mc_, ec_ = [_chart_shape(c) for c in mw._charts], [_chart_shape(c) for c in ew._charts]
        if len(mc_) != len(ec_):
            add("charts", "chart count", len(mc_), len(ec_))
        for i, (x, y) in enumerate(zip(mc_, ec_, strict=False)):
            if x != y:
                add(f"chart {i + 1}", "chart", x, y)

        # row heights (reported)
        off = 0
        for r in range(1, nrows + 1):
            mh = mw.row_dimensions[r].height if r in mw.row_dimensions else None
            eh = ew.row_dimensions[r].height if r in ew.row_dimensions else None
            if (mh or 15) != (eh or 15):
                off += 1
        res.row_heights[tab] = (off, nrows)

    # drawings: anchors, sizes, media
    md, ed = _drawings(master_path, m), _drawings(export_path, e)
    for tab in m_tabs:
        ma, ea = md.get(tab, []), ed.get(tab, [])
        if len(ma) != len(ea):
            res.differences.append(Difference(tab, "drawings", "drawing count", len(ma), len(ea)))
        for i, (x, y) in enumerate(zip(ma, ea, strict=False)):
            where = f"drawing {i + 1}"
            if (x["kind"], x["col"], x["row"], x["sha"]) != (
                y["kind"],
                y["col"],
                y["row"],
                y["sha"],
            ):
                res.differences.append(
                    Difference(
                        tab,
                        where,
                        "anchor",
                        (x["kind"], x["col"], x["row"], x["sha"]),
                        (y["kind"], y["col"], y["row"], y["sha"]),
                    )
                )
            box_a, box_b = (x["x"], x["y"], x["w"], x["h"]), (y["x"], y["y"], y["w"], y["h"])
            if any(abs(p - q) > BOX_TOLERANCE_PX for p, q in zip(box_a, box_b, strict=True)):
                res.differences.append(
                    Difference(
                        tab,
                        where,
                        "box px",
                        tuple(round(v) for v in box_a),
                        tuple(round(v) for v in box_b),
                    )
                )
    return res


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print(__doc__)
        return 2
    result = diff(argv[0], argv[1])
    print(result.report())
    return 1 if result.unexpected() else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
