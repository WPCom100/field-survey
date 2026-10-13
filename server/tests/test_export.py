"""Excel export (T29, PLAN §7.6): TP-4 round trip of the master, the `/exports` routes, the
job lifecycle, the daily export backup, markup flattening and the empty-project workbook."""

import hashlib
import io
import json
import os
import uuid
from datetime import UTC, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
import tp5_precheck
import xlsx_diff
import xlsxwriter
from conftest import HDR
from PIL import Image
from sqlalchemy import func, select, update
from test_seed_import import FIXTURE, import_file, needs_fixture, new_project

from app import projects
from app.auth import system_user, upsert_user
from app.blobs import object_key
from app.config import settings
from app.ids import uuid7
from app.jobs import export_xlsx, exports, markup
from app.models import Blob, Device, Export, Level, Op, PunchItem

# Master quirks TP-4 documents: cells whose master value cannot come back from the records.
ALLOW = {
    # 03-01's Photo cell names photo 2, but photo 2's Device ID names 03-02 only, so the import
    # links photo 2 to 03-02 alone and the export lists it there alone.
    ("Devices", "Photo", ("03", "03-01")): "allow: 03-01 names photo 2, photo 2 names 03-02",
    # Qty belongs to the cable: the master gave B2 C-06 a Qty and left its other end B1 C-05
    # blank; both rows now show the cable's 1.
    ("Connections", "Qty", ("B1", "C-05")): "allow: Qty per cable, the master left one end blank",
}
EXPECTED_DOCUMENTED = {
    "D51": 1,  # no Change Log tab
    "D50": 7,  # the four bus spellings in seven From cells
    # Punch List!Device of a linked item is derived from its device: the device's fuller Room /
    # Location, the rack's and NSPS's model (`Equipment Rack 18U`, `P-NSPS-D-KIT`), `UFMP4` for
    # the master's `UFMP4 #1` (server/docs/exports.md lists them)
    "D148": 87,
    "photo list spacing": 2,  # `1,5` and `12,13`
    "device ID range": 1,  # `01-S8 to 01-S12`
    "Source parenthetical": 3,  # `Connections C-06 (fiber to Level 01)` on three parts
    "allow: 03-01 names photo 2, photo 2 names 03-02": 1,
    "allow: Qty per cable, the master left one end blank": 1,
}


def run_export_via_api(client, factory, store, project_id, params=None, user="Landry"):
    headers = {**HDR, "X-Survey-User": user}
    body = {"project_id": str(project_id)}
    if params is not None:
        body["params"] = params
    r = client.post("/exports", json=body, headers=headers)
    assert r.status_code == 202, r.text
    assert r.json()["job_state"] == "queued" and r.json()["kind"] == "xlsx"
    assert exports.run_queued(factory, store) == ["done"]
    got = client.get(f"/exports/{r.json()['id']}", headers=HDR).json()
    assert got["job_state"] == "done", got["stats"]
    return got


# ------------------------------------------------------------------- TP-4


@needs_fixture
def test_tp4_master_round_trip(session, store, factory, client, tmp_path):
    project_id, batch = import_file(session, store, factory, FIXTURE)
    assert batch.job_state == "done", batch.stats

    got = run_export_via_api(client, factory, store, project_id)
    stats = got["stats"]
    assert got["filename"] == "Lighting_Control_Survey_Master_B2-06_1.xlsx"
    assert stats["rows"] == {
        "Devices": 210,
        "Connections": 107,
        "Punch List": 215,
        "Missing Parts": 85,
        "Photo Log": 76,
    }
    assert stats["images"] == 76 and stats["images_missing"] == 0 and stats["markup_flattened"] == 0
    assert stats["levels"] == ["B2", "B1", "01", "02", "03", "04", "06", "06M"]
    assert got["diff_summary"] is None  # the first export has nothing to compare with

    # the file: a complete blob, downloadable with its filename
    f = client.get(f"/exports/{got['id']}/file", headers=HDR)
    assert f.status_code == 200
    assert f.headers["content-disposition"] == f'attachment; filename="{got["filename"]}"'
    data = f.content
    assert hashlib.sha256(data).hexdigest() == got["blob"]
    blob = session.get(Blob, got["blob"])
    assert (
        blob.upload_state == "complete"
        and blob.mime == exports.XLSX_MIME
        and blob.size == len(data)
    )
    assert client.get(f"/blobs/{got['blob']}/content", headers=HDR).content == data

    out = tmp_path / "export.xlsx"
    out.write_bytes(data)
    # TP-5 (T30): `TP5_OUT=<path> uv run pytest -k tp4` keeps the file to open in Excel.
    if os.environ.get("TP5_OUT"):
        Path(os.environ["TP5_OUT"]).write_bytes(data)
    # TP-5 pre-check (T30): nothing in the package that makes Excel offer a repair
    assert tp5_precheck.check(out) == []
    result = xlsx_diff.diff(FIXTURE, out, allow=ALLOW)
    assert result.unexpected() == [], result.report()
    print("\n".join(str(d) for d in result.differences if d.reason == "D148"))
    assert dict(result.documented()) == EXPECTED_DOCUMENTED, result.report()
    # Row heights are estimated (Excel has no autofit on write): nearly all match the master's.
    for tab, (off, rows) in result.row_heights.items():
        assert off <= max(2, rows * 0.03), (tab, off, rows, result.report())

    # the list route
    listed = client.get("/exports", params={"project_id": str(project_id)}, headers=HDR).json()
    assert [x["id"] for x in listed] == [got["id"]]


def test_tp5_precheck_catches_repair_causes(tmp_path):
    good = tmp_path / "good.xlsx"
    wb = xlsxwriter.Workbook(good)
    ws = wb.add_worksheet("Devices")
    ws.add_table("A3:B5", {"name": "Devices", "columns": [{"header": "Level"}, {"header": "ID"}]})
    wb.close()
    assert tp5_precheck.check(good) == []

    bad = tmp_path / "bad.xlsx"
    wb = xlsxwriter.Workbook(bad)
    ws = wb.add_worksheet("Devices")
    ws.add_table("A3:B5", {"name": "Devices", "columns": [{"header": "Level"}, {"header": "ID"}]})
    ws.write("B3", "ID ")  # header cell no longer equals the table column: Excel repairs this
    ws.set_row(7, 500)  # taller than Excel allows
    wb.close()
    problems = tp5_precheck.check(bad)
    assert any("header 'ID'" in x for x in problems), problems
    assert any("height 500" in x for x in problems), problems


@needs_fixture
def test_second_export_reports_the_diff_and_a_levels_filter(session, store, factory, client):
    project_id, _ = import_file(session, store, factory, FIXTURE)
    first = run_export_via_api(client, factory, store, project_id)

    # an edit between exports (a direct table write stands in for a synced op here)
    b2 = session.scalars(
        select(Level).where(Level.project_id == project_id, Level.level == "B2")
    ).one()
    dev = session.scalars(
        select(Device).where(Device.level_id == b2.id, Device.display_id == "B2-01")
    ).one()
    session.execute(
        update(Device).where(Device.id == dev.id).values(status="Complete", notes="Done")
    )
    session.commit()

    second = run_export_via_api(client, factory, store, project_id)
    assert second["filename"] == "Lighting_Control_Survey_Master_B2-06_2.xlsx"
    ds = second["diff_summary"]
    assert ds["previous"] == first["id"]
    assert ds["tabs"]["Devices"] == {
        "rows": 210,
        "added": 0,
        "removed": 0,
        "changed": 1,
        "examples": ["B2 B2-01"],
    }
    assert ds["tabs"]["Connections"]["changed"] == 0
    # the Summary counts follow the data (cached values Excel shows before recalculating)
    wb = openpyxl.load_workbook(io.BytesIO(store.read(object_key(second["blob"]))))
    assert wb["Devices"]["H4"].value == "Complete"
    assert wb["Summary"]["H8"].value.startswith("=SUMPRODUCT(")
    cached = openpyxl.load_workbook(
        io.BytesIO(store.read(object_key(second["blob"]))), data_only=True
    )
    assert cached["Summary"]["H8"].value == 1 and cached["Summary"]["E8"].value == 2

    # levels filter: just B2
    only = run_export_via_api(client, factory, store, project_id, params={"levels": ["B2"]})
    assert only["stats"]["rows"]["Devices"] == 5 and only["stats"]["levels"] == ["B2"]
    wb = openpyxl.load_workbook(io.BytesIO(store.read(object_key(only["blob"]))))
    assert wb["Devices"]["A2"].value == "Levels B2-B2"
    assert wb["Summary"]["A3"].value == "Surveyed 2026-10-09"
    assert wb["Summary"]["A9"].value == "Total"
    assert len(wb["Summary"]._charts) == 4


@needs_fixture
def test_punch_device_text_follows_the_device(session, store, factory, client):
    """D148: a linked punch item's Device column is derived from its device, so a rename shows
    with no op on the punch item; D150: an unlinked one keeps the master's text."""
    project_id, _ = import_file(session, store, factory, FIXTURE)
    b1, l04 = (
        session.scalars(select(Level).where(Level.project_id == project_id, Level.level == c)).one()
        for c in ("B1", "04")
    )
    dev = session.scalars(
        select(Device).where(Device.level_id == b1.id, Device.display_id == "B1-S3")
    ).one()
    item = session.scalars(
        select(PunchItem).where(PunchItem.level_id == b1.id, PunchItem.number == 8)
    ).one()
    assert item.device_id == dev.id
    assert item.device_text == "PI1002 S3 - Corridor XXB1-C005 north"  # the master's spelling
    # rename the device's room and model (a direct table write stands in for synced ops)
    session.execute(
        update(Device)
        .where(Device.id == dev.id)
        .values(device="PI1004", room_location="Corridor XXB1-C005 north end")
    )
    session.commit()

    got = run_export_via_api(client, factory, store, project_id)
    ws = openpyxl.load_workbook(io.BytesIO(store.read(object_key(got["blob"]))))["Punch List"]
    device_of = {(r[0].value, r[1].value): r[2].value for r in ws.iter_rows(min_row=4)}
    assert device_of[("B1", 8)] == "PI1004 S3 - Corridor XXB1-C005 north end"
    assert session.get(PunchItem, item.id).device_text == "PI1002 S3 - Corridor XXB1-C005 north"
    unlinked = session.scalars(
        select(PunchItem).where(PunchItem.level_id == l04.id, PunchItem.number == 22)
    ).one()
    assert unlinked.device_id is None
    assert device_of[("04", 22)] == "Elec XX04-1E01"  # the room, as the master wrote it
    assert device_of[("B2", 9)] == "All devices"  # the Neuron ID row


# ------------------------------------------------------------- empty project


def test_empty_project_exports_a_valid_workbook(session, store, factory, client):
    project_id = new_project(session, "Empty")
    got = run_export_via_api(client, factory, store, project_id)
    assert got["filename"] == "Lighting_Control_Survey_Master_none_1.xlsx"
    wb = openpyxl.load_workbook(io.BytesIO(store.read(object_key(got["blob"]))))
    assert wb.sheetnames == [
        "Summary",
        "Devices",
        "Connections",
        "Punch List",
        "Missing Parts",
        "Photo Log",
        "Lists",
    ]
    assert wb["Devices"]["A2"].value == "Levels"
    assert (
        set(wb["Devices"].tables) == {"Devices"} and wb["Devices"].tables["Devices"].ref == "A3:K4"
    )
    assert wb["Summary"]["A1"].value == "ETC Lighting Controls Survey"
    assert wb["Summary"]["A3"].value == "Surveyed"
    assert wb["Lists"]["D7"].value == "Not started"
    assert wb["Devices"].freeze_panes == "C4"
    assert len(wb["Summary"]._charts) == 1  # only the by-owner table has rows to chart


def test_tabs_param_writes_only_those_tabs(session, store, factory, client, tmp_path):
    """T31 / D141: `params.tabs` leaves data tabs out; Summary and Lists stay, and no Summary
    formula points at a missing sheet."""
    project_id = new_project(session, "Tabs")
    got = run_export_via_api(
        client, factory, store, project_id, params={"tabs": ["Punch List", "Devices"]}
    )
    assert got["stats"]["tabs"] == ["Devices", "Punch List"]
    assert set(got["stats"]["rows"]) == {"Devices", "Punch List"}
    data = store.read(object_key(got["blob"]))
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Devices", "Punch List", "Lists"]
    formulas = [
        c.value
        for row in wb["Summary"].iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert formulas and not any("Missing Parts" in f for f in formulas)
    out = tmp_path / "tabs.xlsx"
    out.write_bytes(data)
    assert tp5_precheck.check(out) == []
    # a previous export with every tab: the diff compares only the tabs both files have
    again = run_export_via_api(client, factory, store, project_id, params={"tabs": ["Devices"]})
    assert set(again["diff_summary"]["tabs"]) == {"Devices"}

    headers = {**HDR, "X-Survey-User": "Landry"}
    for bad in ([], ["Change Log"], "Devices", ["Summary"]):
        r = client.post(
            "/exports",
            json={"project_id": str(project_id), "params": {"tabs": bad}},
            headers=headers,
        )
        assert r.status_code == 422, bad


# ---------------------------------------------------------------- the API


def test_exports_api_errors(client, session):
    project_id = str(new_project(session))
    headers = {**HDR, "X-Survey-User": "Landry"}
    assert client.post("/exports", json={"project_id": project_id}).status_code == 401
    assert client.post("/exports", json={"project_id": project_id}, headers=HDR).status_code == 400
    assert (
        client.post("/exports", json={"project_id": str(uuid7())}, headers=headers).status_code
        == 404
    )
    r = client.post(
        "/exports", json={"project_id": project_id, "kind": "floor_pdf"}, headers=headers
    )
    assert r.status_code == 422
    r = client.post(
        "/exports", json={"project_id": project_id, "params": {"levels": "B2"}}, headers=headers
    )
    assert r.status_code == 422
    assert client.get(f"/exports/{uuid7()}", headers=HDR).status_code == 404
    queued = client.post("/exports", json={"project_id": project_id}, headers=headers).json()
    assert client.get(f"/exports/{queued['id']}/file", headers=HDR).status_code == 409  # not ready


def test_failed_export_is_recorded(session, store, factory):
    project_id = new_project(session)
    # a floor PDF of a level the project does not have fails cleanly and is recorded
    row = exports.create_export(
        session, project_id=project_id, kind="floor_pdf", params={"level": "ZZ"}, user=None
    )
    assert exports.run_export(factory, store, row.id) == "failed"
    session.expire_all()
    assert "unknown level ZZ" in session.get(Export, row.id).stats["error"]
    assert exports.run_queued(factory, store) == []


# ------------------------------------------------------- daily export backup


def edit_project(session, project_id):
    """One more op on the project (an edit arriving), as the server stores it."""
    seq = session.scalar(select(func.max(Op.server_seq)).where(Op.project_id == project_id)) or 0
    user = upsert_user(session, "Landry")
    session.add(
        Op(
            op_id=uuid7(),
            project_id=project_id,
            user_id=user.id,
            install_id=uuid7(),
            client_ts=datetime.now(UTC),
            server_seq=seq + 1,
            server_ts=datetime.now(UTC),
            entity_kind="project",
            entity_id=project_id,
            kind="set",
            field="name",
            now="Renamed",
            change_type="Field survey",
            state="applied",
        )
    )
    session.commit()


@pytest.fixture
def daily(monkeypatch):
    """Backups on from 00 UTC, the once-a-day memo cleared; returns a 'worker restarted' reset."""
    monkeypatch.setattr(settings, "backup_enabled", True)
    monkeypatch.setattr(settings, "backup_hour_utc", 0)
    monkeypatch.setattr(exports, "_daily_checked_on", None)
    return lambda: monkeypatch.setattr(exports, "_daily_checked_on", None)


def scheduled_rows(session, *project_ids):
    session.expire_all()
    return session.scalars(
        select(Export)
        .where(Export.project_id.in_(project_ids))
        .order_by(Export.created_at, Export.id)
    ).all()


def test_daily_backup_runs_once_a_day_per_project(session, store, factory, monkeypatch, daily):
    p1, p2 = new_project(session, "One"), new_project(session, "Two")
    now = datetime.now(UTC)  # export rows are stamped by the database clock
    monkeypatch.setattr(settings, "backup_hour_utc", 23)
    if now.hour < 23:
        assert exports.daily_tick(factory, store, now) == []  # before BACKUP_HOUR_UTC
    monkeypatch.setattr(settings, "backup_hour_utc", 0)
    assert exports.daily_tick(factory, store, now) == ["done", "done"]
    rows = scheduled_rows(session, p1, p2)
    assert len(rows) == 2
    export_user = system_user(session, "export")
    assert all(r.created_by_id == export_user.id and r.params["scheduled"] for r in rows)
    assert all(isinstance(r.params["server_seq"], int) for r in rows)
    # the same day again: nothing (checked once a day) …
    assert exports.daily_tick(factory, store, now + timedelta(hours=1)) == []
    # … not even after a worker restart: a scheduled export in the last 20 hours
    daily()
    assert exports.daily_tick(factory, store, now + timedelta(hours=1)) == []
    assert len(scheduled_rows(session, p1, p2)) == 2


def test_daily_backup_skips_a_project_with_no_edit(session, store, factory, daily):
    p1, p2 = new_project(session, "One"), new_project(session, "Two")
    now = datetime.now(UTC)
    assert exports.daily_tick(factory, store, now) == ["done", "done"]
    # the next day: nothing changed anywhere → no backup, logged as unchanged
    daily()
    assert exports.daily_tick(factory, store, now + timedelta(days=1)) == []
    # the day after: an edit on One → One alone is backed up
    edit_project(session, p1)
    daily()
    assert exports.daily_tick(factory, store, now + timedelta(days=2)) == ["done"]
    rows = scheduled_rows(session, p1, p2)
    assert [r.project_id for r in rows].count(p1) == 2 and [r.project_id for r in rows].count(
        p2
    ) == 1
    # a backup that missed a picture (still uploading) is taken again the next day
    last_p2 = [r for r in rows if r.project_id == p2][-1]
    session.execute(
        update(Export)
        .where(Export.id == last_p2.id)
        .values(stats={**last_p2.stats, "images_missing": 1})
    )
    session.commit()
    daily()
    assert exports.daily_tick(factory, store, now + timedelta(days=3)) == ["done"]
    assert [r.project_id for r in scheduled_rows(session, p1, p2)].count(p2) == 2


def test_daily_backup_is_off_without_backup_enabled(session, store, factory, monkeypatch, daily):
    p1 = new_project(session, "One")
    monkeypatch.setattr(settings, "backup_enabled", False)
    assert exports.daily_tick(factory, store, datetime.now(UTC)) == []
    assert scheduled_rows(session, p1) == []


def test_backups_older_than_six_months_are_pruned(session, store, factory, client, daily):
    p1 = new_project(session, "One")
    now = datetime.now(UTC)
    assert settings.export_backup_retention_days == 183
    for day in range(4):  # four daily backups, each after an edit
        edit_project(session, p1)
        daily()
        assert exports.daily_tick(factory, store, now + timedelta(days=day)) == ["done"]
    rows = scheduled_rows(session, p1)
    # an empty project exports the same bytes every day: give each backup a file of its own
    for i, r in enumerate(rows):
        data = f"backup {i}".encode()
        sha = hashlib.sha256(data).hexdigest()
        store.put_bytes(object_key(sha), data)
        session.add(
            Blob(sha256=sha, size=len(data), mime=exports.XLSX_MIME, upload_state="complete")
        )
        session.execute(update(Export).where(Export.id == r.id).values(blob=sha))
    session.commit()
    rows = scheduled_rows(session, p1)
    # age the three oldest past six months; the newest stays recent
    for i, r in enumerate(rows[:3]):
        session.execute(
            update(Export).where(Export.id == r.id).values(created_at=now - timedelta(days=200 - i))
        )
    session.commit()
    assert sorted(exports.prune_backups(factory, store, now)) == sorted(str(r.id) for r in rows[:3])
    session.expire_all()
    for r in rows[:3]:
        row = session.get(Export, r.id)
        assert row.stats["pruned_at"] and row.blob == r.blob  # the row stays in the history
        with pytest.raises(KeyError):
            store.read(object_key(r.blob))
        assert session.get(Blob, r.blob) is None
        got = client.get(f"/exports/{r.id}/file", headers=HDR)
        assert got.status_code == 404 and got.json()["detail"]["error"] == "not_available"
    assert store.read(object_key(rows[3].blob))  # the newest is kept
    assert exports.prune_backups(factory, store, now) == []  # nothing twice

    # each project's newest backup is kept however old (no edit for months: still current)
    session.execute(
        update(Export).where(Export.id == rows[3].id).values(created_at=now - timedelta(days=190))
    )
    session.commit()
    assert exports.prune_backups(factory, store, now) == []
    assert store.read(object_key(rows[3].blob))


def test_prune_keeps_a_file_another_export_uses(session, store, factory, daily):
    p1 = new_project(session, "One")
    now = datetime.now(UTC)
    for day in range(2):
        edit_project(session, p1)
        daily()
        assert exports.daily_tick(factory, store, now + timedelta(days=day)) == ["done"]
    old, new = scheduled_rows(session, p1)
    # a manual export that happens to have the same bytes as the old backup
    manual = exports.create_export(session, project_id=p1, params=None, user=None)
    session.execute(
        update(Export).where(Export.id == manual.id).values(job_state="done", blob=old.blob)
    )
    session.execute(
        update(Export).where(Export.id == old.id).values(created_at=now - timedelta(days=200))
    )
    session.commit()
    assert exports.prune_backups(factory, store, now) == [str(old.id)]
    assert store.read(object_key(old.blob))  # still used by the manual export
    assert session.get(Blob, old.blob) is not None


# --------------------------------------------------------------- markup


def jpeg(width=320, height=240, color=(200, 200, 200)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buf, format="JPEG", quality=90)
    return buf.getvalue()


def strokes(*items, width=4032, height=3024) -> bytes:
    return json.dumps(
        {
            "format": "field-survey/strokes",
            "version": 1,
            "space": "image",
            "width": width,
            "height": height,
            "strokes": list(items),
        }
    ).encode()


def test_markup_flatten_draws_the_strokes_and_leaves_unmarked_photos_alone():
    photo = jpeg()
    assert markup.flatten(photo, None) == (photo, 0)
    assert markup.flatten(photo, strokes()) == (photo, 0)  # an empty list changes nothing
    # a red pen stroke across the middle of a 4032-px image lands on the 320-px thumbnail scaled
    pen = {
        "id": "a",
        "tool": "pen",
        "color": "#e53935",
        "width": 60,
        "points": [400, 1512, 0.5, 3600, 1512, 0.5],
        "at": "2026-10-10T14:05:11Z",
    }
    text = {
        "id": "b",
        "tool": "text",
        "color": "#111111",
        "width": 300,
        "text": "LVD",
        "points": [600, 800, 0.5],
        "at": "2026-10-10T14:05:11Z",
    }
    arrow = {
        "id": "c",
        "tool": "arrow",
        "color": "#1e88e5",
        "width": 40,
        "points": [200, 200, 0.5, 1800, 2600, 0.5],
        "at": "x",
    }
    ellipse = {
        "id": "d",
        "tool": "ellipse",
        "color": "#43a047",
        "width": 40,
        "points": [2400, 1800, 0.5, 3800, 2900, 0.5],
        "at": "x",
    }
    high = {
        "id": "e",
        "tool": "highlighter",
        "color": "#fdd835",
        "width": 200,
        "points": [100, 2900, 0.5, 3900, 2900, 0.5],
        "at": "x",
    }
    mark = {
        "id": "f",
        "tool": "marker",
        "color": "#e53935",
        "width": 300,
        "points": [3600, 400, 0.5],
        "ref": {"kind": "photo", "id": "p"},
        "at": "x",
    }
    bad = {"id": "g", "tool": "pen", "points": [1, 2]}  # malformed: skipped, not fatal
    out, skipped = markup.flatten(photo, strokes(pen, text, arrow, ellipse, high, mark, bad))
    assert skipped == 1 and out != photo
    im = Image.open(io.BytesIO(out)).convert("RGB")
    assert im.size == (320, 240)
    r, g, b = im.getpixel((160, 120))  # on the pen line
    assert r > 180 and g < 110 and b < 110
    assert (
        im.getpixel((10, 10)) == (200, 200, 200) or sum(im.getpixel((10, 10))) > 560
    )  # untouched corner
    with pytest.raises(markup.MarkupError):
        markup.flatten(photo, b'{"format": "other", "version": 1, "strokes": []}')
    with pytest.raises(markup.MarkupError):
        markup.flatten(photo, strokes(width=10).replace(b'"version": 1', b'"version": 2'))


def test_export_flattens_markup_onto_the_thumbnail(session, store, factory, client):
    """A photo with a markup blob gets its strokes drawn on the Photo Log picture."""
    from app.sync import OpIn, Push

    user = upsert_user(session, "Landry")
    out = projects.create_project(
        session, projects.ProjectIn(name="Marked", levels=[projects.LevelIn(level="06")]), user
    )
    project_id = uuid.UUID(out.id)
    level = session.scalars(select(Level).where(Level.project_id == project_id)).one()
    thumb = jpeg()
    pen = {
        "id": "a",
        "tool": "pen",
        "color": "#e53935",
        "width": 60,
        "points": [400, 1512, 0.5, 3600, 1512, 0.5],
        "at": "x",
    }
    mk = strokes(pen)
    blobs = {}
    for data, mime in ((thumb, "image/jpeg"), (mk, "application/json")):
        sha = hashlib.sha256(data).hexdigest()
        store.put_file(object_key(sha), io.BytesIO(data), mime)
        session.add(
            Blob(sha256=sha, size=len(data), mime=mime, upload_state="complete", chunks_received=[])
        )
        blobs[mime] = sha
    session.commit()
    photo_id = uuid7()
    op = OpIn(
        op_id=uuid7(),
        project_id=project_id,
        user_id=user.id,
        install_id=uuid7(),
        client_ts=datetime.now(UTC),
        entity_kind="photo",
        entity_id=photo_id,
        kind="create",
        change_type="Field survey",
        reason="Field survey",
        payload={
            "id": str(photo_id),
            "level_id": str(level.id),
            "subject": "ERP",
            "thumb_blob": blobs["image/jpeg"],
            "markup_blob": blobs["application/json"],
        },
    )
    Push(session, exports_body(project_id)).run([op])
    session.commit()

    got = run_export_via_api(client, factory, store, project_id)
    assert got["stats"]["markup_flattened"] == 1 and got["stats"]["images"] == 1
    xlsx = store.read(object_key(got["blob"]))
    import zipfile

    with zipfile.ZipFile(io.BytesIO(xlsx)) as z:
        media = [n for n in z.namelist() if n.startswith("xl/media/")]
        assert len(media) == 1
        im = Image.open(io.BytesIO(z.read(media[0]))).convert("RGB")
    r, g, b = im.getpixel((160, 120))
    assert r > 180 and g < 110 and b < 110
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert wb["Photo Log"]["B4"].value == 1 and wb["Photo Log"].row_dimensions[4].height == 186


def exports_body(project_id):
    from app.jobs.runner import _Body

    return _Body(project_id)


# ----------------------------------------------------- pure helpers


def test_export_helpers():
    assert export_xlsx.number_refs([14, 8, 9, 10, 11, 12, 13]) == "#8-14"
    assert export_xlsx.number_refs([15, 19, 20]) == "#15, #19, #20"
    assert export_xlsx.number_refs([17, 18, 19, 21, 22, 23, 24]) == "#17-19, #21-24"
    assert export_xlsx.number_refs([11, 12]) == "#11, #12"
    ids = ["06M-D1", "06M-S2", "06M-01", "06M-S10", "06M-S9", None]
    assert sorted(ids, key=export_xlsx.id_sort_key) == [
        "06M-01",
        "06M-S2",
        "06M-S9",
        "06M-S10",
        "06M-D1",
        None,
    ]
    assert sorted(["C-10", "C-02", "C-1"], key=export_xlsx.id_sort_key) == ["C-1", "C-02", "C-10"]
    assert export_xlsx.family("06M") == "06" and export_xlsx.family("B2") == "B2"
    assert export_xlsx.wrapped_lines("short", 20) == 1
    assert export_xlsx.wrapped_lines("Outstanding items", 18.43, 14) == 2
    assert export_xlsx.wrapped_lines("Missing Parts", 18.43, 14) == 1
    assert export_xlsx.row_height([("a" * 100, 50)]) == 33.0
    assert exports.filename_for(["B2", "06M"], 4) == "Lighting_Control_Survey_Master_B2-06_4.xlsx"
    dev = Device(display_id="01-S3", device="PI1002", room_location="Corridor XXB1-C005 north")
    assert export_xlsx.default_device_text(dev) == "PI1002 S3 - Corridor XXB1-C005 north"
    assert (
        export_xlsx.default_device_text(Device(display_id="B2-01", device="CTB2LCPB")) == "CTB2LCPB"
    )
    assert export_xlsx.default_device_text(None) is None
    # D151: named panels, racks and NSPS by name alone, an EBDK by its panel, a Foundry panel
    # without a panel name with its room
    room = "Elec XXB1-2E02"
    named = Device(
        display_id="B1-01",
        category="Relay panel",
        device="CTB1LCPB",
        panel_name="CTB1LCPB",
        room_location=room,
    )
    assert export_xlsx.default_device_text(named) == "CTB1LCPB"
    ebdk = Device(
        display_id="B1-02",
        category="EBDK",
        device="EBDK",
        panel_name="CTB1LCPB",
        room_location=room,
    )
    assert export_xlsx.default_device_text(ebdk) == "EBDK (CTB1LCPB)"
    rack = Device(
        display_id="01-05",
        category="Equipment rack",
        device="Equipment Rack 18U",
        room_location=room,
    )
    assert export_xlsx.default_device_text(rack) == "Equipment Rack 18U"
    ufmp = Device(
        display_id="01-08", category="Foundry panel", device="UFMP4", room_location="BH OBS RM 47"
    )
    assert export_xlsx.default_device_text(ufmp) == "UFMP4 - BH OBS RM 47"
